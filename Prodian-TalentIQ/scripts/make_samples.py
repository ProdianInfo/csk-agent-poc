"""Generate fictitious example resumes, including a scanned PDF for OCR tests."""
from pathlib import Path
from io import BytesIO
from docx import Document
import fitz
from PIL import Image,ImageDraw,ImageFont

DIR=Path(__file__).resolve().parents[1]/'samples'/'resumes'
DIR.mkdir(parents=True,exist_ok=True)
text='''SYNTHETIC CANDIDATE PROFILE
Ravi Example
ravi.example@example.invalid
Total Experience: 5 years
Skills: Java, Spring Boot, PostgreSQL, REST API, Azure, Docker
Location: Chennai
This sample contains NO real candidate contact information.
'''
(DIR/'ravi_example.txt').write_text(text)
word=Document();word.add_heading('SYNTHETIC Candidate Profile',0)
for line in text.replace('Ravi Example','Riya Sample').replace('ravi.example@example.invalid','riya.sample@example.invalid').splitlines()[1:]:word.add_paragraph(line)
word.save(DIR/'ravi_example.docx')
pdf=fitz.open();page=pdf.new_page(width=595,height=842)
y=85
for line in text.replace('Ravi Example','Kavin Demo').replace('ravi.example@example.invalid','kavin.demo@example.invalid').splitlines():page.insert_text((65,y),line,fontsize=12);y+=28
pdf.save(DIR/'ravi_example_text.pdf');pdf.close()
img=Image.new('RGB',(1240,1754),'white');draw=ImageDraw.Draw(img)
try:font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',29)
except OSError:font=ImageFont.load_default()
y=90
for line in text.replace('Ravi Example','Meera Mock').replace('ravi.example@example.invalid','meera.mock@example.invalid').splitlines():draw.text((85,y),line,font=font,fill='#202020');y+=64
b=BytesIO();img.save(b,format='PNG')
pdf=fitz.open();page=pdf.new_page(width=595,height=842)
page.insert_image(page.rect,stream=b.getvalue())
pdf.save(DIR/'ravi_example_scanned.pdf');pdf.close()
print('Generated example TXT, DOCX, searchable PDF and scanned PDF.')
