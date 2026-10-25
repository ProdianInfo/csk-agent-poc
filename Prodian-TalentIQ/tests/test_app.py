from io import BytesIO
from docx import Document


def test_page_and_health(client):
    assert client.get('/').status_code==200
    assert client.get('/api/ops/health').json()['status']=='ok'
    assert client.get('/api/candidates').status_code==401

def test_login_roles_and_search(hr):
    r=hr.get('/api/candidates',params={'skills':'Java, Spring Boot','min_years':3})
    assert r.status_code==200
    data=r.json(); assert data['total']>0
    assert data['items'][0]['matched_skills']
    assert data['items'][0]['score']<=100
    assert data['items'][0]['email'].endswith('.invalid')
    assert hr.get('/api/jobs').json()['items']

def test_lead_cannot_see_contact_or_inject(client):
    client.post('/api/auth/login',json={'username':'lead','password':'UnitTest#2026'})
    candidates=client.get('/api/candidates').json()['items']
    assert candidates[0]['email']=='[restricted]'
    assert client.post('/api/ops/inject',json={'kind':'search_500'}).status_code==403

def test_resume_docx_import(hr):
    doc=Document();doc.add_paragraph('Demo Person');doc.add_paragraph('demo.person@example.invalid')
    doc.add_paragraph('Skills: Java, Spring Boot, PostgreSQL, REST API')
    doc.add_paragraph('5 years of experience')
    stream=BytesIO();doc.save(stream)
    r=hr.post('/api/imports/upload',files=[('files',('demo.docx',stream.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'))])
    assert r.status_code==200,r.text
    result=r.json(); assert result['imported']==1,result
    assert result['items'][0]['method']=='docx'
    r2=hr.post('/api/imports/upload',files=[('files',('demo.docx',stream.getvalue(),'application/octet-stream'))])
    assert r2.json()['failed']==1

def test_csv_import(hr):
    text='name,email,years,skills,location\nSynthetic Import,synthetic-import@example.invalid,4,"Python, Azure",Chennai\n'
    r=hr.post('/api/imports/upload',files=[('files',('import.csv',text.encode(),'text/csv'))])
    assert r.json()['imported']==1

def test_incident_injection_and_readonly_tela(admin):
    r=admin.post('/api/ops/inject',json={'kind':'search_500'})
    assert r.status_code==200
    assert r.json()['trace_id']
    assert admin.get('/api/ops/incidents').json()['items']
    assert admin.get('/api/ops/dependency-map').status_code==200
    assert admin.post('/api/tela/tools/search-candidates',json={'skills':'Java'}).status_code==403
    authorized=admin.post('/api/tela/tools/search-candidates',json={'skills':'Java'},headers={'X-TELA-Service-Key':'test-tela-secret'})
    assert authorized.status_code==200
    row=authorized.json()['items'][0]
    assert 'email' not in row and 'phone' not in row and 'display_name' not in row

def test_approval_four_eyes(client):
    client.post('/api/auth/login',json={'username':'hr','password':'UnitTest#2026'})
    r=client.get('/api/ops/incidents').json()['items']
    if not r:
        client.post('/api/auth/login',json={'username':'admin','password':'UnitTest#2026'})
        incident_id=client.post('/api/ops/inject',json={'kind':'schema_drift'}).json()['id']
        client.post('/api/auth/login',json={'username':'hr','password':'UnitTest#2026'})
    else: incident_id=r[0]['id']
    pending=client.post('/api/ops/reprocess/prepare',json={'incident_id':incident_id}).json()
    assert pending['status']=='pending'
    client.post('/api/auth/login',json={'username':'admin','password':'UnitTest#2026'})
    pid=pending['approval_id']
    assert client.post(f'/api/ops/reprocess/{pid}/approve').status_code==200
    assert client.post(f'/api/ops/reprocess/{pid}/execute').json()['status']=='executed'


def test_isolated_real_lab_http_failure(admin):
    response=admin.post('/api/ops/lab/reproduce',json={'kind':'search_500'})
    assert response.status_code==500
    assert response.json()['lab'] is True
    assert response.headers.get('X-Lab-Trace-ID')
    assert admin.get('/api/ops/health').status_code==200
