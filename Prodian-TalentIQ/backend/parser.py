"""Resume extraction for local PDFs, DOCX, scanned PDF OCR, TXT and images.

All extracted fields are *provisional* and must be verified by an HR user.
"""
import io
import re
import tempfile
from pathlib import Path
from .matcher import detect_skills

SUPPORTED = {'.pdf', '.docx', '.doc', '.txt', '.png', '.jpg', '.jpeg', '.tif', '.tiff', '.csv'}

class ParseError(Exception):
    pass


def extract_document(filename, raw):
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ParseError(f'Unsupported file type: {ext}')
    if ext == '.csv':
        return raw.decode('utf-8-sig', errors='replace'), 'csv'
    if ext == '.txt':
        return raw.decode('utf-8-sig', errors='replace'), 'text'
    if ext == '.docx':
        from docx import Document
        d = Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(' | '.join(c.text for c in row.cells))
        return '\n'.join(parts).strip(), 'docx'
    if ext == '.doc':
        import subprocess
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / 'resume.doc'
            source.write_bytes(raw)
            try:
                subprocess.run(['libreoffice', '-env:UserInstallation=file://' + td + '/lo-profile',
                    '--headless', '--convert-to', 'docx', '--outdir', td, str(source)],
                    capture_output=True, timeout=35, check=False)
                converted = Path(td) / 'resume.docx'
                if not converted.exists():
                    raise ParseError('Legacy .doc conversion failed; install LibreOffice or save as DOCX')
                return extract_document('resume.docx', converted.read_bytes())
            except FileNotFoundError:
                raise ParseError('Install LibreOffice for legacy .doc files, or convert to DOCX')
            except subprocess.TimeoutExpired:
                raise ParseError('Legacy .doc conversion timed out')
    if ext == '.pdf':
        from pypdf import PdfReader
        try:
            reader = PdfReader(io.BytesIO(raw), strict=False)
            if len(reader.pages) > 30:
                raise ParseError('PDF exceeds 30-page limit')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError(f'Cannot open PDF: {exc}')
        if len(text.strip()) >= 80:
            return text, 'pdf-text'
        try:
            import fitz, pytesseract
            from PIL import Image
            doc = fitz.open(stream=raw, filetype='pdf')
            ocr = []
            for i in range(min(doc.page_count, 12)):
                page = doc.load_page(i)
                pix = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
                im = Image.open(io.BytesIO(pix.tobytes('png')))
                ocr.append(pytesseract.image_to_string(im, config='--psm 3', timeout=25))
            return '\n'.join(ocr), 'pdf-ocr'
        except Exception as exc:
            raise ParseError('OCR failed. Install Tesseract and PyMuPDF: ' + str(exc))
    try:
        import pytesseract
        from PIL import Image
        image = Image.open(io.BytesIO(raw))
        if image.width * image.height > 25_000_000:
            raise ParseError('Image too large for OCR')
        return pytesseract.image_to_string(image, timeout=30), 'image-ocr'
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError('Image OCR failed: ' + str(exc))


# Prevent qualification headings and resume section titles from becoming names.
NOT_A_NAME = re.compile(
    r'\b(?:resume|cv|curriculum|vitae|profile|summary|education|qualification|skills|'
    r'objective|contact|address|email|mobile|phone|college|university|academic|'
    r'computer\s+science|engineering|engineer|developer|technical|technology|'
    r'professional|experience|employment|career|project|certification|declaration|'
    r'bachelor|master|degree|secondary|higher|percentage|school|personal|details|'
    r'python|java|postgresql|mysql|react|angular|azure|aws|docker|sql|kafka|redis)\b|'
    r'\b(?:b\.?e\.?|b\.?tech\.?|m\.?e\.?|m\.?tech\.?|bsc|msc|bca|mca|mba)\b', re.I)
GENERIC_FILE_WORDS = re.compile(r'\b(?:resume|cv|curriculum|vitae|profile|copy|final|latest|updated|doc|document|scan|scanned|pdf|candidate)\b', re.I)
EMAIL = re.compile(r'\b[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}\b', re.I)
PHONE = re.compile(r'(?<!\d)(?:\+91[\s\-]?)?[6-9]\d{4}[\s\-]?\d{5}(?!\d)')
NAME_LABEL = re.compile(r'^(?:full\s+name|candidate\s+name|name)\s*[:\-]\s*(.+)$', re.I)
NAME_SHAPE = re.compile(r"^[A-Za-z][A-Za-z.' -]{2,65}$")


def _looks_like_name(value):
    value = ' '.join(value.strip().split()).strip(' .-')
    return (3 <= len(value) <= 65 and len(value.split()) <= 5
            and NAME_SHAPE.fullmatch(value) is not None
            and NOT_A_NAME.search(value) is None)


def _name_from_filename(filename):
    stem = Path(filename).stem
    stem = re.sub(r'[_\-]+', ' ', stem)
    stem = GENERIC_FILE_WORDS.sub(' ', stem)
    stem = re.sub(r'\b\d{2,}\b', ' ', stem)
    stem = ' '.join(stem.split()).strip()
    return stem if _looks_like_name(stem) else ''


def _name_from_email(address):
    if not address:
        return ''
    local = address.split('@')[0].split('+')[0]
    if re.match(r'^(?:info|jobs|hr|admin|recruitment|recruiter|contact|resume|cv|hello)\b', local, re.I):
        return ''
    local = re.sub(r'\d+$', '', local)
    local = re.sub(r'[._-]+', ' ', local).strip()
    return local if _looks_like_name(local) else ''


def infer_name(lines, filename, email):
    """Prefer explicit names, then early header, useful filename, then email.

    This is an *unverified suggestion*; not an identity validation mechanism.
    """
    for line in lines[:24]:
        labelled = NAME_LABEL.match(line)
        if labelled and _looks_like_name(labelled.group(1)):
            return labelled.group(1).strip(), 'explicit name field'
    # A genuine header name is usually in the first few lines. Do not mistake
    # middle-of-resume lines such as degree names and headings for identities.
    for line in lines[:4]:
        if _looks_like_name(line):
            return line, 'resume header'
    from_file = _name_from_filename(filename)
    if from_file:
        return from_file, 'filename'
    from_email = _name_from_email(email)
    if from_email:
        return from_email, 'email address'
    return 'Name Requires Review', 'unavailable'


def extract_experience(text):
    """Return fractional years only for explicitly stated experience; otherwise unknown.

    Does not infer experience from employment dates because overlapping jobs,
    internships, and education dates make unverified calculations misleading.
    """
    patterns = [
        r'\b(?:total|overall|professional|relevant|work|industry|it)\s+(?:work\s+)?experience\s*[:\-]?\s*'
        r'(\d{1,2}(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\b(?:\s*(?:and|&)\s*(\d{1,2})\s*months?)?',
        r'\bexperience\s*[:\-]\s*(\d{1,2}(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\b'
        r'(?:\s*(?:and|&)\s*(\d{1,2})\s*months?)?',
        r'\b(\d{1,2}(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\b'
        r'(?:\s*(?:and|&)\s*(\d{1,2})\s*months?)?\s+(?:of\s+)?(?:overall\s+|total\s+|professional\s+|relevant\s+|work\s+)?(?:experience|exp\b)',
    ]
    for pattern in patterns:
        m = re.search(pattern, text, re.I)
        if m:
            years = float(m.group(1))
            months = int(m.group(2)) if m.lastindex and m.group(2) else 0
            if 0 <= years <= 60 and 0 <= months <= 11:
                return round(years + months / 12, 2)
    return None


def parse_candidate(filename, raw):
    text, method = extract_document(filename, raw)
    if not text.strip():
        raise ParseError('No readable text detected; manual review required')
    lines = [' '.join(s.strip().split()) for s in text.splitlines() if s.strip()]
    email_match = EMAIL.search(text)
    phone_match = PHONE.search(text)
    email = email_match.group() if email_match else ''
    name, name_source = infer_name(lines, filename, email)
    years = extract_experience(text)
    warnings = []
    if name_source != 'explicit name field':
        warnings.append('Name inferred from ' + name_source + '; verify against resume')
    if not email_match:
        warnings.append('Email missing')
    if not phone_match:
        warnings.append('Phone missing')
    if years is None:
        warnings.append('Experience not found; verify manually (0 is not a verified value)')
    return {'name':name.title(), 'email':email,
        'phone':re.sub(r'\s+', '', phone_match.group()) if phone_match else '',
        'years':years if years is not None else 0, 'skills':detect_skills(text),
        'extracted_text':text[:100000], 'parse_method':method, 'review_needed':True,
        'warnings':warnings}
